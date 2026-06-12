import json
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

# Read post numbers from new-list.txt
with open("new-list.txt", "r", encoding="utf-8") as f:
    post_numbers = []
    for line in f:
        line = line.strip()
        if line:
            post_numbers.append(int(line))

# Read voters.json
with open("data/voters.json", "r", encoding="utf-8") as f:
    voters = json.load(f)

# Build lookup by post_number
voter_lookup = {}
for voter in voters:
    pn = voter.get("post_number")
    if pn is not None:
        if pn not in voter_lookup:
            voter_lookup[pn] = voter
        else:
            # If duplicate, prefer existing or life_member over new
            existing = voter_lookup[pn]
            pref = {"new": 0, "existing": 1, "life_member": 2}
            if pref.get(voter.get("renewal", ""), 0) > pref.get(existing.get("renewal", ""), 0):
                voter_lookup[pn] = voter

# Create workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "New Entries"

# Header style
header_font = Font(bold=True, size=11)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font_white = Font(bold=True, size=11, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Write headers
headers = ["Serial Number", "Name", "License Number", "Renewal Status", "Phone Number", "Source", "Address"]
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.font = header_font_white
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center')
    cell.border = thin_border

# Write data
row_num = 2
not_found = []
for idx, post_num in enumerate(post_numbers, 1):
    voter = voter_lookup.get(post_num)
    if voter:
        source_raw = voter.get("source", "")
        if source_raw == "Life Member":
            source = "life"
        else:
            source = "voter"

        renewal = voter.get("renewal", "")

        values = [
            idx,
            voter.get("name", ""),
            voter.get("post_number", ""),
            renewal,
            voter.get("phone", ""),
            source,
            voter.get("address", "")
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.border = thin_border
        row_num += 1
    else:
        not_found.append(post_num)

# Auto-fit column widths
for col in ws.columns:
    max_length = 0
    col_letter = col[0].column_letter
    for cell in col:
        try:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        except:
            pass
    ws.column_dimensions[col_letter].width = min(max_length + 4, 40)

# Write not found post numbers to a separate sheet if any
if not_found:
    ws2 = wb.create_sheet("Not Found")
    ws2.cell(row=1, column=1, value="Post Number").font = header_font_white
    ws2.cell(row=1, column=1).fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    ws2.cell(row=1, column=1).border = thin_border
    for i, pn in enumerate(not_found, 2):
        ws2.cell(row=i, column=1, value=pn).border = thin_border
    ws2.column_dimensions['A'].width = 15

# Save
output_path = "new_entries.xlsx"
wb.save(output_path)
print(f"Excel file saved as: {output_path}")
print(f"Total entries matched: {row_num - 2}")
print(f"Entries not found: {len(not_found)}")
if not_found:
    print(f"Not found post numbers: {not_found[:20]}{'...' if len(not_found) > 20 else ''}")
