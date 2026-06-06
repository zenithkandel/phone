import sys, io, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

wb = openpyxl.load_workbook('C:/xampp/htdocs/codes/phone/DLA_FINAL_VOTER_With_Phones.xlsx')
voters = []
vid = 1

# Sheet 1: Final Voter List
ws1 = wb['Final Voter List']
for i, row in enumerate(ws1.iter_rows(values_only=True)):
    if i == 0:
        continue
    serial, receipt_no, name, post_number, title, renewal, signature, phone = row
    voters.append({
        "id": vid,
        "serial": serial,
        "receipt_no": receipt_no,
        "name": name if name else "",
        "post_number": post_number,
        "title": title if title else "",
        "renewal": "existing" if renewal == 1 else "new",
        "phone": phone if phone else "Not Found",
        "source": "Final Voter List"
    })
    vid += 1

# Sheet 2: Life Member
ws2 = wb['Life Member']
for i, row in enumerate(ws2.iter_rows(values_only=True)):
    if i == 0:
        continue
    serial, name, title, cert_no, contact, phone = row
    voters.append({
        "id": vid,
        "serial": serial,
        "receipt_no": cert_no,
        "name": name if name else "",
        "post_number": None,
        "title": title if title else "",
        "renewal": "life_member",
        "phone": phone if phone else "Not Found",
        "source": "Life Member"
    })
    vid += 1

with open('C:/xampp/htdocs/codes/phone/data/voters.json', 'w', encoding='utf-8') as f:
    json.dump(voters, f, ensure_ascii=False, indent=2)

print(f"Converted {len(voters)} voters to JSON")
print(f"  Final Voter List: {ws1.max_row - 1}")
print(f"  Life Members: {ws2.max_row - 1}")
